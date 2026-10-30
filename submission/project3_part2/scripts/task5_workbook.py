#!/usr/bin/env python3
"""Write the Task 5 presentation workbook without changing its source workbook.

Primary calculations are supplied by analyze_task5.py. This module only writes
the presentation tables and checks that existing worksheets survive unchanged.
Requires openpyxl; a local /tmp/puf_excel_deps installation is also supported.
"""

from copy import copy
from pathlib import Path
import hashlib
import sys


try:
    from openpyxl import load_workbook
except ModuleNotFoundError:
    dependency_path = Path("/tmp/puf_excel_deps")
    if dependency_path.is_dir():
        sys.path.insert(0, str(dependency_path))
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as error:
        raise ImportError(
            "Task 5 Excel export requires openpyxl. Install it with "
            "python3 -m pip install openpyxl, then rerun the analysis."
        ) from error

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter


TEXT = "263B49"
BLUE = "DCEAF4"
GREEN = "E8F3EC"
SAND = "FFF5DB"
STRIPE = "F5F8FA"
LINE = "D7E1E8"
MUTED = "556D7B"
NUMBER = "0.000000"
PERCENT = "0.000000%"


def _title(worksheet, title, subtitle, last_column):
    last_letter = get_column_letter(last_column)
    worksheet.merge_cells(f"A1:{last_letter}1")
    cell = worksheet["A1"]
    cell.value = title
    cell.font = Font(name="Calibri", size=17, bold=True, color=TEXT)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 31
    worksheet.merge_cells(f"A2:{last_letter}2")
    cell = worksheet["A2"]
    cell.value = subtitle
    cell.font = Font(name="Calibri", size=11, color=MUTED)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[2].height = 32
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 95
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins = PageMargins(
        left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.2, footer=0.2
    )
    worksheet.oddFooter.right.text = "Page &P of &N"


def _section(worksheet, cell_range, title, fill=BLUE):
    worksheet.merge_cells(cell_range)
    cell = worksheet[cell_range.split(":")[0]]
    cell.value = title
    cell.font = Font(name="Calibri", size=11, bold=True, color=TEXT)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[cell.row].height = 24


def _metric(worksheet, row, label, value, label_range, value_range, number_format=None):
    worksheet.merge_cells(f"{label_range[0]}{row}:{label_range[1]}{row}")
    worksheet.merge_cells(f"{value_range[0]}{row}:{value_range[1]}{row}")
    label_cell = worksheet[f"{label_range[0]}{row}"]
    label_cell.value = label
    label_cell.font = Font(name="Calibri", size=11, color=TEXT)
    label_cell.alignment = Alignment(wrap_text=True, vertical="center")
    value_cell = worksheet[f"{value_range[0]}{row}"]
    value_cell.value = value
    value_cell.font = Font(name="Calibri", size=11, bold=True, color=TEXT)
    value_cell.alignment = Alignment(vertical="center", horizontal="left")
    if number_format:
        value_cell.number_format = number_format
    worksheet.row_dimensions[row].height = 23


def _headers(worksheet, row, values, start=1):
    for offset, value in enumerate(values):
        cell = worksheet.cell(row, start + offset, value)
        cell.font = Font(name="Calibri", size=11, bold=True, color=TEXT)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color=LINE))
    worksheet.row_dimensions[row].height = 33


def _table(worksheet, title, subtitle, headers, data, widths, text_columns=(),
           decimal_columns=(), percent_columns=()):
    _title(worksheet, title, subtitle, len(headers))
    _headers(worksheet, 4, headers)
    text_columns = set(text_columns)
    decimal_columns = set(decimal_columns)
    percent_columns = set(percent_columns)
    normal_font = Font(name="Calibri", size=11, color=TEXT)
    binary_font = Font(name="Consolas", size=11, color=TEXT)
    stripe = PatternFill("solid", fgColor=STRIPE)
    for index, values in enumerate(data, 5):
        for column, value in enumerate(values, 1):
            cell = worksheet.cell(index, column, value)
            cell.font = binary_font if column in text_columns else normal_font
            cell.alignment = Alignment(vertical="center", horizontal="center")
            if index % 2:
                cell.fill = stripe
            if column in text_columns:
                if value is not None and not isinstance(value, str):
                    raise TypeError(f"Text column contains a non-string at {cell.coordinate}")
                cell.number_format = "@"
            elif column in percent_columns:
                cell.number_format = PERCENT
            elif column in decimal_columns:
                cell.number_format = NUMBER
    last_letter = get_column_letter(len(headers))
    worksheet.auto_filter.ref = f"A4:{last_letter}{max(4, worksheet.max_row)}"
    worksheet.freeze_panes = "A5"
    worksheet.print_title_rows = "1:4"
    worksheet.print_area = f"A1:{last_letter}{worksheet.max_row}"
    for column, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(column)].width = width


def _summary(worksheet, summary):
    _title(
        worksheet, "Part 2 / Task 5 — RO PUF hardware results",
        "Different-challenge variation uses recorded valid CRPs. Repeatability "
        "is separately labeled as user-reported evidence; raw repeat rows are unavailable.", 12,
    )
    for letter in "ABCDEFGHIJKL":
        worksheet.column_dimensions[letter].width = 13
    worksheet.freeze_panes = "A5"
    counts = summary["counts"]
    variation = summary["variation"]
    repeatability = summary["repeatability"]
    _section(worksheet, "A4:F4", "Collection and validity")
    for row, label, value in [
        (5, "Recorded challenges", counts["total_challenges"]),
        (6, "Valid challenges included", counts["valid_challenges"]),
        (7, "Invalid challenges excluded", counts["excluded_challenges"]),
        (8, "Same-index exclusions", counts["same_index_excluded"]),
        (9, "Different-index exclusions", counts["different_index_excluded"]),
        (10, "Blank invalid masks confirmed 0000", counts["blank_invalid_confirmed_zero"]),
        (11, "Explicit 0000 invalid masks", counts["explicit_zero_masks"]),
    ]:
        _metric(worksheet, row, label, value, ("A", "D"), ("E", "F"))

    _section(worksheet, "H4:L4", "Different-challenge variation")
    for row, label, value, fmt in [
        (5, "Unordered challenge pairs", variation["comparisons"], None),
        (6, "Total mismatched bits", variation["total_differing_bits"], None),
        (7, "Total compared bits", variation["total_compared_bits"], None),
        (8, "Mean HD (bits / 4)", variation["mean_bits"], NUMBER),
        (9, "Mean normalized HD", variation["percent"] / 100, PERCENT),
        (10, "Unbiased-response benchmark", variation["benchmark_percent"] / 100, "0%"),
        (11, "Minimum / maximum HD (bits)", f"{variation['min_bits']} / {variation['max_bits']}", None),
    ]:
        _metric(worksheet, row, label, value, ("H", "J"), ("K", "L"), fmt)
    worksheet.row_dimensions[10].height = 31

    _section(worksheet, "A14:L14", "Same-challenge repeatability — user-reported observations", GREEN)
    _metric(worksheet, 15, "Fresh-reset trials per challenge", repeatability["reported_trials_per_challenge"], ("A", "D"), ("E", "F"))
    _metric(worksheet, 16, "Challenges reported unchanged", repeatability["reported_unchanged_challenges"], ("A", "D"), ("E", "F"))
    _metric(worksheet, 15, "Reported repeatability HD", repeatability["reported_hd_percent"] / 100, ("H", "J"), ("K", "L"), "0%")
    _metric(worksheet, 16, "Repeatability ideal", 0, ("H", "J"), ("K", "L"), "0%")
    _metric(worksheet, 17, "Raw trial records available", "Yes" if repeatability["raw_trial_records_available"] else "No", ("A", "D"), ("E", "F"))
    _metric(worksheet, 17, "Computed raw-repeat pairs", repeatability["computed_raw_repeat_comparisons"], ("H", "J"), ("K", "L"))
    worksheet.merge_cells("A19:L21")
    worksheet["A19"] = repeatability["evidence_basis"]
    worksheet["A19"].font = Font(name="Calibri", size=11, color=TEXT)
    worksheet["A19"].fill = PatternFill("solid", fgColor=SAND)
    worksheet["A19"].alignment = Alignment(wrap_text=True, vertical="center")

    _section(worksheet, "A24:D24", "Pairwise HD distribution")
    _headers(worksheet, 25, ["HD (bits)", "HD (%)", "Pair count", "Share of pairs"])
    for index, entry in enumerate(variation["distribution"], 26):
        values = [entry["hd_bits"], entry["hd_percent"] / 100,
                  entry["pairs"], entry["pair_fraction"]]
        for column, value in enumerate(values, 1):
            cell = worksheet.cell(index, column, value)
            cell.font = Font(name="Calibri", size=11, color=TEXT)
            cell.alignment = Alignment(horizontal="center")
            if column in (2, 4):
                cell.number_format = PERCENT

    _section(worksheet, "F24:I24", "Response-bit balance (valid CRPs)")
    _headers(worksheet, 25, ["Bit", "Zeros", "Ones", "Ones (%)"], 6)
    for index, entry in enumerate(summary["bit_balance"], 26):
        for column, value in enumerate([
            entry["bit_label"], entry["zeros"], entry["ones"], entry["ones_fraction"]
        ], 6):
            cell = worksheet.cell(index, column, value)
            cell.font = Font(name="Calibri", size=11, color=TEXT)
            cell.alignment = Alignment(horizontal="center")
            if column == 9:
                cell.number_format = PERCENT

    _section(worksheet, "A34:D34", "Response frequency (valid CRPs)")
    _headers(worksheet, 35, ["Response", "Count", "Share", ""], 1)
    for index, entry in enumerate(summary["response_distribution"], 36):
        worksheet.cell(index, 1, entry["response_bin"]).number_format = "@"
        worksheet.cell(index, 2, entry["count"])
        worksheet.cell(index, 3, entry["fraction"]).number_format = PERCENT
        for cell in worksheet[index][:3]:
            cell.font = Font(name="Consolas" if cell.column == 1 else "Calibri", size=11, color=TEXT)
            cell.alignment = Alignment(horizontal="center")

    _section(worksheet, "F34:L34", "Provenance and interpretation")
    notes = [
        f"Source: {summary['provenance']['source_file']}",
        f"SHA-256: {summary['provenance']['source_sha256']}",
        summary["provenance"]["method"],
        *summary["provenance"]["confirmations"],
        "The original CRPs sheet is preserved, including its historical comments. "
        "This summary records the later user clarification of blank invalid masks and repeat coverage.",
        "A nonzero invalid mask excludes the complete four-bit CRP; no tied or invalid bits are silently masked.",
        "Fixed-condition repeatability does not establish reliability under different temperatures, supplies, or boards.",
    ]
    for index, note in enumerate(notes):
        row = 36 + 2 * index
        worksheet.merge_cells(start_row=row, start_column=6, end_row=row + 1, end_column=12)
        cell = worksheet.cell(row, 6, note)
        cell.font = Font(name="Calibri", size=10, color=MUTED)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        worksheet.row_dimensions[row].height = 20
        worksheet.row_dimensions[row + 1].height = 20
    worksheet.print_area = f"A1:L{max(52, 36 + 2 * len(notes))}"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToHeight = 1


def _sheet_snapshot(worksheet):
    return {
        "cells": {
            cell.coordinate: (
                cell.value, cell.data_type, copy(cell.font), copy(cell.fill),
                copy(cell.border), copy(cell.alignment), copy(cell.protection),
                cell.number_format,
                None if cell.comment is None else (cell.comment.text, cell.comment.author),
            )
            for row in worksheet for cell in row
        },
        "merges": tuple(str(value) for value in worksheet.merged_cells.ranges),
        "validations": ET_string(worksheet.data_validations.to_tree()),
        "freeze_panes": worksheet.freeze_panes,
        "auto_filter": worksheet.auto_filter.ref,
        "dimensions": (worksheet.max_row, worksheet.max_column),
    }


def ET_string(element):
    # openpyxl may use either lxml or ElementTree, so use its own serializer.
    from openpyxl.xml.functions import tostring
    return tostring(element)


def write_results_workbook(source_path, output_path, summary, rows, pairs, by_challenge):
    """Create a separate workbook with copied CRPs and four Task 5 result sheets."""
    source_path, output_path = Path(source_path), Path(output_path)
    if source_path.resolve() == output_path.resolve():
        raise ValueError("Task 5 output must be separate from the source workbook.")
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    workbook = load_workbook(source_path)
    original_sheets = list(workbook.sheetnames)
    snapshots = {sheet.title: _sheet_snapshot(sheet) for sheet in workbook}
    output_names = ["Task5_Summary", "Task5_By_Challenge", "Task5_HD_Pairs", "Task5_Exclusions"]
    if set(output_names) & set(original_sheets):
        raise ValueError("Source already contains Task 5 sheets; refusing to replace existing sheets.")

    _summary(workbook.create_sheet(output_names[0]), summary)
    _table(
        workbook.create_sheet(output_names[1]), "Valid CRPs: HD against other challenges",
        "Each row is compared with every other valid challenge on the same physical FPGA.",
        ["Challenge (Cha1||Cha0)", "Cha1 (SW7–4)", "Cha0 (SW3–0)", "Response (LEDG3–0)",
         "Comparisons", "Total HD (bits)", "Mean HD (bits)", "Mean HD (%)", "Min HD (bits)", "Max HD (bits)"],
        [[r["challenge_bin"], r["cha1_bin"], r["cha0_bin"], r["response_bin"],
          r["comparisons"], r["total_differing_bits"], r["mean_bits"], r["percent"] / 100,
          r["min_bits"], r["max_bits"]] for r in by_challenge],
        [25, 17, 17, 23, 16, 19, 20, 20, 17, 17],
        text_columns=(1, 2, 3, 4), decimal_columns=(7,), percent_columns=(8,),
    )
    _table(
        workbook.create_sheet(output_names[2]), "All unordered valid challenge pairs",
        "HD is the popcount of the four-bit response XOR; normalized HD = HD / 4.",
        ["Challenge A", "Cha1 A", "Cha0 A", "Response A", "Challenge B", "Cha1 B", "Cha0 B",
         "Response B", "Response XOR", "HD (bits)", "HD (%)"],
        [[r["challenge_a_bin"], r["cha1_a_bin"], r["cha0_a_bin"], r["response_a_bin"],
          r["challenge_b_bin"], r["cha1_b_bin"], r["cha0_b_bin"], r["response_b_bin"],
          r["xor_bin"], r["hd_bits"], r["hd_percent"] / 100] for r in pairs],
        [18, 12, 12, 15, 18, 12, 12, 15, 17, 13, 18],
        text_columns=tuple(range(1, 10)), percent_columns=(11,),
    )
    excluded = [r for r in rows if not r["included"]]
    _table(
        workbook.create_sheet(output_names[3]), "Excluded hardware CRPs",
        "All original measurements remain in CRPs. Any nonzero invalid flag excludes the full response from the primary calculation.",
        ["Source row", "Challenge", "Cha1", "Cha0", "Response", "Raw invalid", "Effective invalid",
         "Done", "Exclusion reason", "Status evidence"],
        [[r["source_row"], r["challenge_bin"], r["cha1_bin"], r["cha0_bin"], r["response_bin"],
          r["raw_invalid_bin"], r["invalid_bin"], str(r["done"]), r["exclusion_reason"], r["status_basis"]]
         for r in excluded],
        [14, 18, 12, 12, 15, 16, 18, 10, 45, 65],
        text_columns=(2, 3, 4, 5, 6, 7, 8, 9, 10),
    )
    workbook.active = workbook.sheetnames.index("Task5_Summary")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    check = load_workbook(output_path)
    try:
        if check.sheetnames != original_sheets + output_names:
            raise RuntimeError("Task 5 workbook sheet preservation check failed.")
        for name, expected in snapshots.items():
            if _sheet_snapshot(check[name]) != expected:
                raise RuntimeError(f"Original sheet changed during export: {name}")
        if check["Task5_HD_Pairs"].max_row != len(pairs) + 4:
            raise RuntimeError("Pairwise workbook table row count does not match the supplied analysis.")
        if check["Task5_By_Challenge"].max_row != len(by_challenge) + 4:
            raise RuntimeError("Per-challenge workbook table row count is incorrect.")
        if check["Task5_Exclusions"].max_row != len(excluded) + 4:
            raise RuntimeError("Exclusions workbook table row count is incorrect.")
    finally:
        check.close()
    if hashlib.sha256(source_path.read_bytes()).hexdigest() != source_hash:
        raise RuntimeError("Source workbook changed during Task 5 export.")
    return output_path
