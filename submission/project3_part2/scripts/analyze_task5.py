#!/usr/bin/env python3
"""Reproduce Part 2 Task 5 from the completed, source-hashed hardware workbook.

Run: python3 project3_part2/scripts/analyze_task5.py
Requires openpyxl; also submit analyze_crps.py and task5_workbook.py.
The original workbook is never edited. Collection clarifications below apply
only to the confirmed input hash; review them if the source workbook changes.
"""

import argparse
from collections import Counter
import csv
import hashlib
import itertools
import json
from pathlib import Path
import re
import shutil
import sys
import tempfile


DEPENDENCIES = Path("/tmp/puf_excel_deps")
if DEPENDENCIES.is_dir():
    sys.path.insert(0, str(DEPENDENCIES))
try:
    from openpyxl import load_workbook
except ImportError as error:
    raise SystemExit("Install the Excel dependency with: python3 -m pip install openpyxl") from error

from analyze_crps import analyze as analyze_hardware_csv


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "tb/ro_puf_256_crps.xlsx"
DEFAULT_OUTPUT = ROOT / "tb/task5_results"
CONFIRMED_SOURCE_SHA256 = "31cb68a766893f65d01c25427342866157cbc8b8a7dc70602a3a4031aac6c7ed"
HEADERS = ("Cha1", "Cha0", "Challenge (combined)", "Response", "Invalid LEDs (LEDR3–0)")
CONFIRMATIONS = [
    "User confirmed every blank invalid-mask cell means LEDR3..0=0000 after Done became high.",
    "User reported all 256 challenges were measured after five independent resets with unchanged responses.",
    "The workbook stores one response per challenge, not five individual raw trial records.",
]
RESPONSE_PATTERN = re.compile(r"([01]{4})(?:\s*\(0[xX]([0-9a-fA-F])\))?")
ROW_FIELDS = [
    "device_id", "source_row", "cha0_bin", "cha1_bin", "challenge_bin", "repeat",
    "record_type", "response_bin", "raw_invalid_bin", "invalid_bin", "done",
    "included", "exclusion_reason", "status_basis",
]
PAIR_FIELDS = [
    "challenge_a_bin", "cha1_a_bin", "cha0_a_bin", "response_a_bin",
    "challenge_b_bin", "cha1_b_bin", "cha0_b_bin", "response_b_bin",
    "xor_bin", "hd_bits", "hd_percent",
]
CHALLENGE_FIELDS = [
    "challenge_bin", "cha1_bin", "cha0_bin", "response_bin", "comparisons",
    "total_differing_bits", "mean_bits", "percent", "min_bits", "max_bits",
]


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def four_bits(value, field):
    text = str(value).strip() if value is not None else ""
    if not re.fullmatch(r"[01]{4}", text):
        raise ValueError(f"{field}: expected four binary digits, got {value!r}")
    return text


def response_bits(value):
    match = RESPONSE_PATTERN.fullmatch(str(value).strip())
    if not match:
        raise ValueError(f"Malformed four-bit response: {value!r}")
    binary, hexadecimal = match.groups()
    if hexadecimal is not None and int(binary, 2) != int(hexadecimal, 16):
        raise ValueError(f"Binary and hexadecimal response disagree: {value!r}")
    return binary


def normalize_rows(raw_rows, device_id="FPGA_0"):
    """Normalize confirmed collection semantics without constructing raw repeats."""
    rows = []
    seen = set()
    for source_row, values in enumerate(raw_rows, 2):
        if all(value is None for value in values):
            continue
        cha1, cha0, combined, response, invalid = values[:5]
        cha1 = four_bits(cha1, f"Row {source_row} Cha1")
        cha0 = four_bits(cha0, f"Row {source_row} Cha0")
        if combined != cha1 + cha0:
            raise ValueError(f"Row {source_row}: combined challenge must be Cha1||Cha0")
        if combined in seen:
            raise ValueError(f"Row {source_row}: duplicate challenge {combined}")
        seen.add(combined)
        raw_invalid = str(invalid).strip() if invalid is not None else ""
        mask = four_bits(raw_invalid, f"Row {source_row} invalid mask") if raw_invalid else "0000"
        included = mask == "0000"
        reason = "" if included else (
            "nonzero invalid mask; same RO index" if cha1 == cha0
            else "nonzero invalid mask; different RO indices"
        )
        if included:
            basis = "blank mask and Done confirmed by user" if not raw_invalid else "explicit zero mask; Done confirmed by user"
        else:
            basis = "explicit invalid flag; excluded regardless of unrecorded overall Done"
        rows.append({
            "device_id": device_id,
            "source_row": source_row,
            "cha0_bin": cha0,
            "cha1_bin": cha1,
            "challenge_bin": combined,
            "repeat": 1,
            "record_type": "one canonical workbook record; not an individual trial log",
            "response_bin": response_bits(response),
            "raw_invalid_bin": raw_invalid,
            "invalid_bin": mask,
            "done": "1" if included else "",
            "included": included,
            "exclusion_reason": reason,
            "status_basis": basis,
        })
    return sorted(rows, key=lambda row: row["challenge_bin"])


def read_collection(path, device_id):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook["CRPs"]
        if tuple(cell.value for cell in sheet[1][:5]) != HEADERS:
            raise ValueError("Unexpected CRPs sheet columns")
        return normalize_rows(list(sheet.iter_rows(min_row=2, max_col=5, values_only=True)), device_id)
    finally:
        workbook.close()


def calculate(rows, source, source_hash):
    valid = [row for row in rows if row["included"]]
    excluded = [row for row in rows if not row["included"]]
    n = len(valid)
    if n < 2:
        raise ValueError("At least two valid challenges are required for pairwise HD")
    by_challenge = {
        row["challenge_bin"]: {
            **{key: row[key] for key in CHALLENGE_FIELDS[:4]},
            "comparisons": 0, "total_differing_bits": 0, "min_bits": 4, "max_bits": 0,
        } for row in valid
    }
    pairs = []
    distribution = Counter()
    for a, b in itertools.combinations(valid, 2):
        xor = int(a["response_bin"], 2) ^ int(b["response_bin"], 2)
        distance = xor.bit_count()
        pairs.append({
            "challenge_a_bin": a["challenge_bin"],
            "cha1_a_bin": a["cha1_bin"], "cha0_a_bin": a["cha0_bin"],
            "response_a_bin": a["response_bin"],
            "challenge_b_bin": b["challenge_bin"],
            "cha1_b_bin": b["cha1_bin"], "cha0_b_bin": b["cha0_bin"],
            "response_b_bin": b["response_bin"],
            "xor_bin": f"{xor:04b}", "hd_bits": distance, "hd_percent": distance * 25,
        })
        distribution[distance] += 1
        for record in (a, b):
            per = by_challenge[record["challenge_bin"]]
            per["comparisons"] += 1
            per["total_differing_bits"] += distance
            per["min_bits"] = min(per["min_bits"], distance)
            per["max_bits"] = max(per["max_bits"], distance)
    for per in by_challenge.values():
        per["mean_bits"] = per["total_differing_bits"] / per["comparisons"]
        per["percent"] = per["mean_bits"] * 25
    comparison_count = n * (n - 1) // 2
    total = sum(pair["hd_bits"] for pair in pairs)
    bit_balance = []
    for bit in range(3, -1, -1):
        ones = sum((int(row["response_bin"], 2) >> bit) & 1 for row in valid)
        bit_balance.append({"bit_label": f"Response[{bit}]", "ones": ones,
                            "zeros": n - ones, "ones_fraction": ones / n})
    # Independent identity: for each bit, each 1/0 unordered pair differs once.
    independent_total = sum(bit["ones"] * bit["zeros"] for bit in bit_balance)
    assert total == independent_total
    assert len(pairs) == comparison_count == sum(distribution.values())
    assert sum(per["total_differing_bits"] for per in by_challenge.values()) == 2 * total
    assert all(per["comparisons"] == n - 1 for per in by_challenge.values())
    frequency = Counter(row["response_bin"] for row in valid)
    summary = {
        "response_width": 4,
        "device_id": rows[0]["device_id"],
        "counts": {
            "total_challenges": len(rows), "valid_challenges": n,
            "excluded_challenges": len(excluded),
            "same_index_excluded": sum(row["cha0_bin"] == row["cha1_bin"] for row in excluded),
            "different_index_excluded": sum(row["cha0_bin"] != row["cha1_bin"] for row in excluded),
            "blank_invalid_confirmed_zero": sum(not row["raw_invalid_bin"] for row in rows),
            "explicit_zero_masks": sum(row["raw_invalid_bin"] == "0000" for row in rows),
            "valid_fraction": n / len(rows),
        },
        "variation": {
            "comparisons": comparison_count, "total_differing_bits": total,
            "total_compared_bits": 4 * comparison_count,
            "mean_bits": total / comparison_count,
            "percent": 100 * total / (4 * comparison_count),
            "min_bits": min(distribution), "max_bits": max(distribution),
            "distribution": [{"hd_bits": d, "hd_percent": d * 25,
                              "pairs": distribution[d], "pair_fraction": distribution[d] / comparison_count}
                             for d in range(5)],
            "independent_bit_count_total": independent_total,
            "benchmark_percent": 50,
        },
        "bit_balance": bit_balance,
        "response_distribution": [{"response_bin": f"{value:04b}", "count": frequency[f"{value:04b}"],
                                   "fraction": frequency[f"{value:04b}"] / n}
                                  for value in range(16)],
        "repeatability": {
            "reported_trials_per_challenge": 5,
            "reported_unchanged_challenges": len(rows),
            "reported_hd_bits": 0, "reported_hd_percent": 0,
            "raw_trial_records_available": False,
            "computed_raw_repeat_comparisons": 0,
            "raw_calculated_hd_percent": None,
            "ideal_percent": 0,
            "evidence_basis": CONFIRMATIONS[1] + " " + CONFIRMATIONS[2],
        },
        "provenance": {
            "source_file": str(source.resolve()), "source_sha256": source_hash,
            "confirmations": CONFIRMATIONS,
            "method": "All unordered pairs of valid, distinct challenges; XOR/popcount across four actual response bits, normalized by four.",
            "challenge_order": "SW7..SW0 = Cha1||Cha0; response = LEDG3..LEDG0",
            "blank_invalid_policy": "Normalize to 0000 only under the user's source-specific confirmation; original cells remain unchanged.",
            "done_policy": "User-confirmed high for included measurements; excluded-row completion remains unrecorded.",
            "canonical_record_policy": "One recorded response per challenge, repeat=1 for CSV compatibility; no synthetic repeated measurements.",
        },
    }
    return summary, pairs, list(by_challenge.values())


def write_csv(path, fields, records):
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)


def report_text(summary):
    c, v, repeat = (summary[key] for key in ("counts", "variation", "repeatability"))
    most_biased = max(summary["bit_balance"], key=lambda bit: abs(bit["ones_fraction"] - 0.5))
    difference = abs(v["percent"] - 50)
    bit_lines = "\n".join(
        f"{bit['bit_label']:<14} {bit['ones']:>4} {bit['zeros']:>5} {100 * bit['ones_fraction']:>10.4f}%"
        for bit in summary["bit_balance"]
    )
    distance_lines = "\n".join(
        f"{entry['hd_bits']:>2} bits ({entry['hd_percent']:>3}%): {entry['pairs']:>6} pairs"
        for entry in v["distribution"]
    )
    return f"""PART 2, TASK 5: HARDWARE CRP RESULTS

We configured the FPGA with Project3.sof and collected all {c['total_challenges']} eight-bit
challenge combinations. Cha1 controls SW7..SW4 and Cha0 controls SW3..SW0;
the combined challenge is Cha1||Cha0. Responses are read as LEDG3..LEDG0
after completion. The collection procedure resets the PUF for each trial.

The source worksheet contains {c['total_challenges']} responses. We excluded {c['excluded_challenges']} challenges
with a nonzero Invalid mask: {c['same_index_excluded']} selected the same oscillator index for
both counters, and {c['different_index_excluded']} selected different indices. This leaves {c['valid_challenges']} valid
CRPs, exceeding the required minimum of 16. The valid fraction is
{100 * c['valid_fraction']:.6f}% across the entire challenge space. During collection,
all four red LEDs were off after Done became high for the
{c['blank_invalid_confirmed_zero']} rows whose Invalid cells were left blank. These rows were therefore
treated as having Invalid=0000; {c['explicit_zero_masks']} additional row explicitly records 0000.

SINGLE-CHIP HAMMING DISTANCE

We compared every unordered pair of distinct valid challenges on this
FPGA. For four-bit responses Ra and Rb:
    HD_bits(Ra,Rb) = popcount(Ra XOR Rb)
    HD_percent(Ra,Rb) = 100 * HD_bits(Ra,Rb) / 4

There are C({c['valid_challenges']},2) = {v['comparisons']:,} comparisons, covering {v['total_compared_bits']:,} bit positions.
The comparisons contain {v['total_differing_bits']:,} differing bits. Therefore:
    Mean HD = {v['total_differing_bits']:,} / {v['comparisons']:,} = {v['mean_bits']:.10f} bits out of 4
    Normalized mean HD = {v['percent']:.10f}%

Individual-pair distribution:
{distance_lines}

The measured average is {difference:.6f} percentage points from the 50% benchmark
for independent, unbiased responses to different challenges. It therefore
shows near-ideal average challenge-to-challenge variation in this collected
data. Each individual four-bit comparison is restricted to 0%, 25%, 50%,
75%, or 100%; the reported mean averages all {v['comparisons']:,} comparisons.

The average alone does not establish complete response balance or
independence. Bit balance across the {c['valid_challenges']} valid responses is:
Bit            Ones Zeros        Ones %
{bit_lines}

{most_biased['bit_label']} is the most biased bit, with {most_biased['ones']} ones among {c['valid_challenges']} responses
({100 * most_biased['ones_fraction']:.4f}%). All sixteen four-bit output patterns occur, but their
frequencies differ. These observations qualify the favorable mean HD;
the result does not by itself prove entropy, security, or device uniqueness.
Pairwise comparisons also share responses and are not independent samples.

REPEATABILITY

We repeated every one of the {repeat['reported_unchanged_challenges']} challenges five times, using a fresh
reset for each measurement. Each challenge produced the same response in
all five trials. For the valid CRPs, the observed same-challenge HD was
therefore 0 bits, or 0%, which is the ideal for repeatability. The worksheet
stores only one canonical response per challenge. This repeatability result
is based on our collection observations; individual raw trial records are
unavailable for independent recalculation. No synthetic repeat rows were
created. The {c['excluded_challenges']} flagged invalid challenges remain excluded from valid-CRP
analysis even when their responses repeat.

These measurements support stable responses under the tested conditions.
Temperature, supply-voltage changes, and long-term drift were not
characterized by this collection. Thus the data supports observed
repeatability, rather than a guarantee across operating conditions.

The 50% benchmark applies to responses under DIFFERENT challenges;
the 0% repeatability ideal applies to repeated use of the SAME challenge.

Files: ro_puf_task5_results.xlsx, task5_summary.json, normalized_crps.csv,
valid_crps.csv, excluded_crps.csv, pairwise_hd.csv, by_challenge_hd.csv,
bit_balance.csv, response_distribution.csv, and analysis_checks.txt.
Source SHA256: {summary['provenance']['source_sha256']}
"""


def readme_text():
    return """REPRODUCING PART 2 TASK 5

Run from the Project3 directory:
    python3 project3_part2/scripts/analyze_task5.py

To preserve an existing results directory, choose a fresh location:
    python3 project3_part2/scripts/analyze_task5.py --output-dir /tmp/task5_recheck

Dependencies: Python 3 and openpyxl. The script can use the session-local
/tmp/puf_excel_deps installation, or install openpyxl in your environment:
    python3 -m pip install openpyxl

Submit analyze_task5.py, analyze_crps.py, and task5_workbook.py together,
the completed source workbook, and these analysis outputs. The original
workbook is never edited. Source-specific collection clarifications are
embedded with the confirmed source SHA256; if the workbook changes, review
the clarifications and update the confirmed hash before running again.

normalized_crps.csv contains all 256 canonical source records.
valid_crps.csv is the existing analyze_crps.py-compatible input containing
only confirmed-valid records. excluded_crps.csv preserves all flagged rows;
their overall Done state is unknown and remains blank.

repeat=1 is a compatibility identifier for the ONE stored response per
challenge. It is not a fabrication of the five verbally reported trials.
The existing helper correctly reports no raw repeat comparisons. The
user-reported five unchanged trials and 0% observed repeatability HD are
clearly separated from calculations on stored individual measurements.

CSV binary columns must be imported as Text in spreadsheet tools to retain
leading zeros. The Excel results workbook already stores these as Text.
"""


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--device-id", default="FPGA_0")
    args = parser.parse_args()
    source = args.input.resolve()
    output = args.output_dir.resolve()
    if output.exists() and any(output.iterdir()):
        raise FileExistsError(f"Preserving existing results. Choose a new --output-dir: {output}")
    source_hash = digest(source)
    if source_hash != CONFIRMED_SOURCE_SHA256:
        raise ValueError("Source hash changed. Review the collection confirmations before updating CONFIRMED_SOURCE_SHA256.")
    rows = read_collection(source, args.device_id)
    summary, pairs, by_challenge = calculate(rows, source, source_hash)
    assert summary["counts"]["total_challenges"] == 256
    assert summary["counts"]["valid_challenges"] == 239
    assert summary["counts"]["excluded_challenges"] == 17
    assert summary["variation"]["comparisons"] == 28441
    valid = [row for row in rows if row["included"]]
    excluded = [row for row in rows if not row["included"]]
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".task5-", dir=output.parent) as temporary:
        staging = Path(temporary)
        write_csv(staging / "normalized_crps.csv", ROW_FIELDS, rows)
        write_csv(staging / "valid_crps.csv", ROW_FIELDS, valid)
        write_csv(staging / "excluded_crps.csv", ROW_FIELDS, excluded)
        write_csv(staging / "pairwise_hd.csv", PAIR_FIELDS, pairs)
        write_csv(staging / "by_challenge_hd.csv", CHALLENGE_FIELDS, by_challenge)
        write_csv(staging / "bit_balance.csv", ["bit_label", "ones", "zeros", "ones_fraction"], summary["bit_balance"])
        write_csv(staging / "response_distribution.csv", ["response_bin", "count", "fraction"], summary["response_distribution"])
        helper_result = analyze_hardware_csv(staging / "valid_crps.csv")
        helper_device, = helper_result["devices"]
        for name, expected in {
            "valid_challenges": len(valid), "variation_comparisons": len(pairs),
            "variation_mean_bits": summary["variation"]["mean_bits"],
            "variation_percent": summary["variation"]["percent"],
            "repeatability_comparisons": 0,
        }.items():
            observed = helper_device[name]
            if isinstance(expected, float):
                assert abs(observed - expected) < 1e-12
            else:
                assert observed == expected
        assert helper_device["repeatability_percent"] is None
        summary["existing_csv_helper_check"] = helper_result
        (staging / "task5_summary.json").write_text(json.dumps(summary, indent=4) + "\n", encoding="utf-8")
        (staging / "task5_report_answer.txt").write_text(report_text(summary), encoding="utf-8")
        (staging / "README.txt").write_text(readme_text(), encoding="utf-8")
        from task5_workbook import write_results_workbook
        write_results_workbook(source, staging / "ro_puf_task5_results.xlsx", summary, rows, pairs, by_challenge)
        assert digest(source) == source_hash, "Source workbook changed during analysis"
        checks = [
            "PASS: 256 unique eight-bit challenges and well-formed four-bit responses/masks",
            "PASS: binary and hexadecimal response labels agree",
            "PASS: combined challenge is Cha1||Cha0 with leading zeros preserved",
            "PASS: 239 included and 17 flagged rows excluded (16 same-index, 1 different-index)",
            "PASS: all 28441 unordered pairs represented exactly once",
            "PASS: XOR/popcount total agrees with independent per-bit ones*zeros total",
            "PASS: per-challenge totals count every pair twice, once for each endpoint",
            "PASS: existing analyze_crps.py agrees with valid-record calculations",
            "PASS: no synthetic raw repeats; user-reported repeatability is separately identified",
            "PASS: separate workbook exported and original source file SHA256 unchanged",
        ]
        (staging / "analysis_checks.txt").write_text("\n".join(checks) + "\n", encoding="utf-8")
        output.mkdir(exist_ok=True)
        for artifact in staging.iterdir():
            shutil.copyfile(artifact, output / artifact.name)
    print(f"Task 5 complete: {len(valid)} valid CRPs, {len(excluded)} excluded.")
    print(f"Single-chip variation: {summary['variation']['mean_bits']:.10f} bits / {summary['variation']['percent']:.10f}%.")
    print("Repeatability: user-reported five unchanged trials; 0% observed HD; raw repeats unavailable.")
    print(f"Results: {output}")


if __name__ == "__main__":
    main()
