#!/usr/bin/env python3
"""Create the 256-challenge RO PUF worksheet without overwriting collected data."""

import argparse
from pathlib import Path
import sys


DEPENDENCY_PATH = Path("/tmp/puf_excel_deps")
if DEPENDENCY_PATH.is_dir():
    sys.path.insert(0, str(DEPENDENCY_PATH))

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins


DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / "tb" / "ro_puf_256_crps.xlsx"
HEADERS = (
    "Cha1",
    "Cha0",
    "Challenge (combined)",
    "Response",
    "Invalid LEDs (LEDR3–0)",
)
MEASURED = {
    ("0100", "1011"): "0000",
    ("1111", "1101"): "0001",
    ("1000", "0110"): "0010",
    ("1111", "0110"): "0011",
    ("1000", "0111"): "0100",
    ("0001", "0000"): "0101",
    ("0001", "0010"): "0110",
    ("1011", "0110"): "0111",
    ("1010", "1110"): "1000",
    ("0010", "0001"): "1001",
    ("0000", "0001"): "1010",
    ("0010", "0100"): "1011",
    ("0110", "1100"): "1100",
    ("0110", "1000"): "1101",
    ("1001", "1110"): "1110",
    ("0111", "1110"): "1111",
}
MEASURED_COLOR = "FFE2F0D9"


def response_label(binary):
    return f"{binary} (0x{int(binary, 2):X})"


def create_workbook(destination):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "CRPs"
    workbook.properties.title = "RO PUF: 256 challenge combinations"
    workbook.properties.subject = "Manual hardware CRP collection"
    workbook.properties.description = (
        "All Cha1/Cha0 combinations, with 16 user-supplied responses highlighted "
        "in light green. The user reported five unchanged measurements per "
        "prefilled response. Invalid LED states were not supplied."
    )

    worksheet.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="FF234E70")
    measured_fill = PatternFill("solid", fgColor=MEASURED_COLOR)
    bottom_border = Border(bottom=Side(style="hair", color="FFD9E1E8"))
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 33

    header_comments = {
        "A1": "Cha1 is SW7–SW4. Hold reset while setting the challenge.",
        "B1": "Cha0 is SW3–SW0. Release reset after setting the challenge.",
        "C1": "Eight-bit combined challenge = Cha1 followed by Cha0 (Cha1||Cha0).",
        "D1": (
            "Read response bits from LEDG3 down to LEDG0 after the green Done "
            "LEDG4 is on and the outputs have settled. Format: 1010 (0xA). "
            "Binary-only entries are also allowed. Light-green rows contain "
            "the 16 user-supplied measurements."
        ),
        "E1": (
            "Record red LED states from LEDR3 down to LEDR0 as four binary "
            "digits, after green Done LEDG4 is on and the outputs have settled. "
            "Blank means unrecorded, including for the 16 prefilled responses; "
            "their invalid LED states were not supplied. A 1 marks an invalid "
            "response bit; discard a response if any invalid LED is on."
        ),
    }
    for address, content in header_comments.items():
        worksheet[address].comment = Comment(content, "Collection notes")

    for combined in range(256):
        cha1 = f"{combined >> 4:04b}"
        cha0 = f"{combined & 15:04b}"
        response = MEASURED.get((cha1, cha0))
        worksheet.append([
            cha1,
            cha0,
            f"{combined:08b}",
            response_label(response) if response is not None else None,
            None,
        ])
        row_number = combined + 2
        worksheet.row_dimensions[row_number].height = 21
        for cell in worksheet[row_number]:
            cell.number_format = "@"
            cell.font = Font(name="Consolas", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = bottom_border
            if response is not None:
                cell.fill = measured_fill
        if response is not None:
            worksheet.cell(row_number, 4).comment = Comment(
                "Transcribed from the user's supplied RO PUF CRP image. "
                "The user reported measuring this challenge five times with "
                "no response changes. Individual raw repeat records and "
                "invalid LED states were not supplied.",
                "Measurement provenance",
            )

    response_choices = ",".join(response_label(f"{value:04b}") for value in range(16))
    response_validation = DataValidation(
        type="list", formula1=f'"{response_choices}"', allow_blank=True
    )
    response_validation.showErrorMessage = False
    response_validation.showInputMessage = True
    response_validation.promptTitle = "Response: LEDG3–0"
    response_validation.prompt = "Wait for Done LEDG4. Choose a value or type four binary bits."
    response_validation.showDropDown = False
    worksheet.add_data_validation(response_validation)
    response_validation.add("D2:D257")

    invalid_choices = ",".join(f"{value:04b}" for value in range(16))
    invalid_validation = DataValidation(
        type="list", formula1=f'"{invalid_choices}"', allow_blank=True
    )
    invalid_validation.showErrorMessage = True
    invalid_validation.errorTitle = "Use four binary digits"
    invalid_validation.error = "Enter LEDR3–0 as four bits, such as 0000 or 0101."
    invalid_validation.errorStyle = "stop"
    invalid_validation.showInputMessage = True
    invalid_validation.promptTitle = "Invalid LEDs: LEDR3–0"
    invalid_validation.prompt = "Blank = unrecorded. Record all four red LED states after Done."
    invalid_validation.showDropDown = False
    worksheet.add_data_validation(invalid_validation)
    invalid_validation.add("E2:E257")

    for column, width in {"A": 12, "B": 12, "C": 25, "D": 23, "E": 28}.items():
        worksheet.column_dimensions[column].width = width
    worksheet.auto_filter.ref = "A1:E257"
    worksheet.freeze_panes = "D2"
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 100
    worksheet.print_title_rows = "1:1"
    worksheet.print_area = "A1:E257"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins = PageMargins(
        left=0.3, right=0.3, top=0.45, bottom=0.45, header=0.2, footer=0.2
    )
    worksheet.oddHeader.center.text = "RO PUF — Cha1 / Cha0 CRP Collection"
    worksheet.oddFooter.left.text = "Light green: supplied measurements"
    worksheet.oddFooter.right.text = "Page &P of &N"
    workbook.save(destination)


def verify_workbook(destination):
    workbook = load_workbook(destination)
    assert workbook.sheetnames == ["CRPs"]
    worksheet = workbook["CRPs"]
    assert (worksheet.max_row, worksheet.max_column) == (257, 5)
    assert tuple(cell.value for cell in worksheet[1]) == HEADERS
    observed = {}
    combined_values = []
    for row in worksheet.iter_rows(min_row=2, max_row=257):
        cha1, cha0, combined, response, invalid = [cell.value for cell in row]
        assert isinstance(cha1, str) and len(cha1) == 4
        assert isinstance(cha0, str) and len(cha0) == 4
        assert isinstance(combined, str) and len(combined) == 8
        assert combined == cha1 + cha0
        assert all(cell.data_type == "s" for cell in row[:3])
        assert all(cell.number_format == "@" for cell in row)
        combined_values.append(int(combined, 2))
        key = (cha1, cha0)
        assert key not in observed
        observed[key] = response
        assert invalid is None
        if key in MEASURED:
            assert response == response_label(MEASURED[key])
            assert row[3].comment is not None
            assert "five times" in row[3].comment.text
            assert all(cell.fill.patternType == "solid" for cell in row)
            assert all(cell.fill.fgColor.rgb == MEASURED_COLOR for cell in row)
        else:
            assert response is None
            assert all(cell.fill.patternType is None for cell in row)
    assert combined_values == list(range(256))
    assert len(observed) == 256
    assert sum(value is not None for value in observed.values()) == 16
    assert worksheet.freeze_panes == "D2"
    assert worksheet.auto_filter.ref == "A1:E257"
    assert len(worksheet.data_validations.dataValidation) == 2
    assert "Blank means unrecorded" in worksheet["E1"].comment.text
    workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    if args.output.exists() and not args.overwrite:
        parser.error("Output already exists; choose a new --output to preserve manual measurements.")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    create_workbook(args.output)
    verify_workbook(args.output)
    print(f"Created and verified {args.output}")
    print("256 ordered challenges; 16 correct green measured rows; 240 blank responses; all invalid masks blank.")


if __name__ == "__main__":
    main()
